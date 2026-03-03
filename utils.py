import base64
import streamlit as st
import PyPDF2
import docx
import pandas as pd
import io
from fpdf import FPDF

def load_css(file_name):
    """Loads a CSS file and injects it into the Streamlit app."""
    try:
        with open(file_name) as f:
            st.markdown(f'<style>{f.read()}</style>', unsafe_allow_html=True)
    except Exception as e:
        st.error(f"Failed to load CSS: {e}")

def extract_text(file):
    """Extracts text from PDF, DOCX, or TXT files."""
    if file is None:
        return ""
    
    file_type = file.name.split('.')[-1].lower()
    text = ""
    
    try:
        if file_type == 'txt':
            text = file.getvalue().decode('utf-8')
        elif file_type == 'pdf':
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        elif file_type == 'docx':
            doc = docx.Document(file)
            for para in doc.paragraphs:
                text += para.text + "\n"
        else:
            st.error(f"Unsupported file type: {file_type}")
    except Exception as e:
        st.error(f"Error reading file {file.name}: {e}")
        
    return text

def calculate_performance(json_data):
    """
    Calculates the performance score and bands based on JSON AI output.
    Returns score, band_name, risk_index, supervision_level, html_badge
    """
    if "Evaluation" not in json_data:
         return 0, "Error", "High", "High", ""

    evaluations = json_data["Evaluation"]
    
    total_score = 0
    max_score = len(evaluations) * 100 if len(evaluations) > 0 else 700

    for item in evaluations:
        item_score = 0
        status = str(item.get("Status", "")).strip().upper()
        evidence = str(item.get("Evidence Strength", "")).strip().upper()
        risk = str(item.get("Compliance Risk", "")).strip().upper()
        
        # Base status score (Max 50)
        if status == "YES":
            item_score += 50
            
        # Evidence strength (Max 30)
        if evidence == "STRONG":
            item_score += 30
        elif evidence == "MODERATE":
            item_score += 15
        elif evidence == "WEAK":
            item_score += 5
            
        # Compliance Risk (Max 20) - Low risk is good
        if risk == "LOW":
            item_score += 20
        elif risk == "MEDIUM":
            item_score += 10
            
        total_score += float(item_score)
        
    score = (float(total_score) / float(max_score)) * 100.0 if max_score > 0 else 0.0
    score = round(score, 1)

    if score >= 90:
        band = "Operationally Excellent"
        badge_class = "badge-excellent"
        risk = "Low"
        supervision = "Minimal"
    elif score >= 75:
        band = "Operationally Strong"
        badge_class = "badge-strong"
        risk = "Low-Medium"
        supervision = "Standard"
    elif score >= 50:
        band = "Operationally Moderate"
        badge_class = "badge-moderate"
        risk = "Medium-High"
        supervision = "Close"
    else:
        band = "Operational Risk"
        badge_class = "badge-risk"
        risk = "Critical"
        supervision = "Direct"

    badge_html = f'<div class="badge {badge_class}">{band}</div>'
    
    return score, band, risk, supervision, badge_html

def create_excel_download(df):
    """Converts a DataFrame to an Excel binary object for downloading."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
        df.to_excel(writer, index=False, sheet_name='Audit Report')
        workbook = writer.book
        worksheet = writer.sheets['Audit Report']
        
        # Add some basic styling to Excel
        header_format = workbook.add_format({
            'bold': True,
            'text_wrap': True,
            'valign': 'top',
            'fg_color': '#00122e',
            'font_color': 'white',
            'border': 1
        })
        for col_num, value in enumerate(df.columns.values):
            worksheet.write(0, col_num, value, header_format)
            worksheet.set_column(col_num, col_num, 15)
            
    processed_data = output.getvalue()
    return processed_data

def apply_color_coding(val):
    """Pandas styler wrapper to color code cells"""
    val_str = str(val).strip().upper()
    if val_str == "YES" or val_str == "LOW" or val_str == "STRONG":
        color = '#00ffaa'
    elif val_str == "NO" or val_str == "HIGH" or val_str == "NONE" or val_str == "WEAK" or val_str == "CRITICAL":
        color = '#ff3366'
    elif val_str == "MEDIUM" or val_str == "MODERATE":
        color = '#ffc400'
    else:
        color = 'white'
    return f'color: {color}; font-weight: bold;'

def create_pdf_download(json_data):
    """Creates a PDF report from the evaluation JSON data."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    # Title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Operational Audit Report", ln=True, align='C')
    pdf.ln(10)
    
    if "Executive Summary" in json_data:
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(200, 10, txt="Executive Summary", ln=True)
        pdf.set_font("Arial", size=11)
        for k, v in json_data["Executive Summary"].items():
            # Handle unicode gracefully
            text = f"{k}: {v}".encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 8, txt=text)
        pdf.ln(5)
        
    if "Evaluation" in json_data:
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(200, 10, txt="Audit Details", ln=True)
        pdf.set_font("Arial", size=10)
        for item in json_data["Evaluation"]:
            crit = str(item.get("Criterion", ""))
            status = str(item.get("Status", ""))
            pdf.set_font("Arial", 'B', 10)
            tit_text = f"Criterion: {crit} - Status: {status}".encode('latin-1', 'replace').decode('latin-1')
            pdf.cell(0, 8, txt=tit_text, ln=True)
            pdf.set_font("Arial", size=10)
            for k, v in item.items():
                if k not in ["Criterion", "Status"]:
                    t_text = f"  {k}: {v}".encode('latin-1', 'replace').decode('latin-1')
                    pdf.multi_cell(0, 6, txt=t_text)
            pdf.ln(3)

    return pdf.output(dest='S').encode('latin-1')
